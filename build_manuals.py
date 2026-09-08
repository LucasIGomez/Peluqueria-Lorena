# -*- coding: utf-8 -*-
"""
Script generador de los dos manuales de usuario en formato Word (.docx):
1. Manual de Usuario- Gomez, Ramos.docx (Módulos de Usuarios y Servicios)
2. Manual de Usuario- Mansilla, Chiacchio.docx (Módulos de Inventarios y Proveedores)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Colores institucionales
COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D - Títulos principales
COLOR_GOLD = RGBColor(197, 155, 39)    # #C59B27 - Acento dorado
COLOR_SLATE = RGBColor(74, 85, 104)    # #4A5568 - Subtítulos
COLOR_DARK = RGBColor(30, 41, 59)      # #1E293B - Texto general
COLOR_RED = RGBColor(185, 28, 28)      # #B91C1C - Alertas críticas
COLOR_GREEN = RGBColor(22, 101, 52)    # #166534 - Éxito / Tips

HEX_NAVY = "1B365D"
HEX_GOLD = "C59B27"
HEX_LIGHT_GRAY = "F8FAFC"
HEX_LIGHT_GOLD = "FEF9C3"
HEX_LIGHT_RED = "FEE2E2"
HEX_LIGHT_GREEN = "DCFCE7"

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, box_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)

    if box_type == "tip":
        bg_hex = HEX_LIGHT_GREEN
        border_hex = "166534"
        icon_title = f"💡 TIP DE LORENA: {title}"
        title_color = COLOR_GREEN
    elif box_type == "warning":
        bg_hex = HEX_LIGHT_GOLD
        border_hex = "CA8A04"
        icon_title = f"⚠️ IMPORTANTE: {title}"
        title_color = RGBColor(161, 98, 7)
    elif box_type == "danger":
        bg_hex = HEX_LIGHT_RED
        border_hex = "DC2626"
        icon_title = f"🛑 ATENCIÓN: {title}"
        title_color = COLOR_RED
    else:
        bg_hex = HEX_LIGHT_GRAY
        border_hex = HEX_NAVY
        icon_title = f"ℹ️ INFORMACIÓN: {title}"
        title_color = COLOR_NAVY

    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)

    borders_xml = (
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(icon_title)
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = title_color

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    run_b = p2.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_heading_1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = COLOR_NAVY
    return h

def format_heading_2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12.5)
    r.bold = True
    r.font.color.rgb = COLOR_GOLD
    return h

def format_heading_3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.bold = True
    r.font.color.rgb = COLOR_SLATE
    return h

def add_body_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.bold = True
        rb.font.name = "Arial"
        rb.font.size = Pt(10)
        rb.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK
    return p

def add_step_p(doc, step_num, title, description):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_num = p.add_run(f"Paso {step_num}: ")
    r_num.bold = True
    r_num.font.name = "Arial"
    r_num.font.size = Pt(10)
    r_num.font.color.rgb = COLOR_NAVY

    r_title = p.add_run(f"{title}. ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = COLOR_DARK

    r_desc = p.add_run(description)
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = COLOR_DARK
    return p

def add_bullet_p(doc, title, description):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = COLOR_DARK

    r_desc = p.add_run(description)
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = COLOR_DARK
    return p

def add_styled_table(doc, headers, data_rows, col_widths=None):
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], HEX_NAVY)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, row_data in enumerate(data_rows):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = HEX_LIGHT_GRAY if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_DARK

    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl

def setup_page_layout(doc, title_header):
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = f"Peluquería Lorena — {title_header}"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.runs[0].font.name = "Arial"
        p_hdr.runs[0].font.size = Pt(8.5)
        p_hdr.runs[0].font.color.rgb = COLOR_SLATE

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.text = "Sistema de Gestión Integral • Manual de Uso Práctico"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ftr.runs[0].font.name = "Arial"
        p_ftr.runs[0].font.size = Pt(8.5)
        p_ftr.runs[0].font.color.rgb = COLOR_SLATE

def add_cover_page(doc, main_title, subtitle, authors, module_badge):
    p_salon = doc.add_paragraph()
    p_salon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_salon.paragraph_format.space_before = Pt(70)
    p_salon.paragraph_format.space_after = Pt(4)
    r_salon = p_salon.add_run("PELUQUERÍA LORENA")
    r_salon.font.name = "Arial"
    r_salon.font.size = Pt(26)
    r_salon.bold = True
    r_salon.font.color.rgb = COLOR_GOLD

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_after = Pt(36)
    r_tag = p_tag.add_run("SISTEMA DE GESTIÓN Y ADMINISTRACIÓN OPERATIVA")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.bold = True
    r_tag.font.color.rgb = COLOR_SLATE

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.0)
    set_cell_background(c, HEX_LIGHT_GRAY)
    set_cell_margins(c, top=200, bottom=200, left=200, right=200)

    p_badge = c.paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b = p_badge.add_run(f"★ {module_badge.upper()} ★\n\n")
    r_b.font.name = "Arial"
    r_b.font.size = Pt(11)
    r_b.bold = True
    r_b.font.color.rgb = COLOR_GOLD

    r_title = p_badge.add_run(main_title + "\n\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    r_sub = p_badge.add_run(subtitle)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_SLATE

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(110)
    p_meta.paragraph_format.space_after = Pt(4)
    r_meta1 = p_meta.add_run("Elaborado por: ")
    r_meta1.font.name = "Arial"
    r_meta1.font.size = Pt(11)
    r_meta1.font.color.rgb = COLOR_SLATE
    r_meta2 = p_meta.add_run(authors + "\n")
    r_meta2.font.name = "Arial"
    r_meta2.font.size = Pt(12)
    r_meta2.bold = True
    r_meta2.font.color.rgb = COLOR_NAVY

    p_meta3 = doc.add_paragraph()
    p_meta3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta3 = p_meta3.add_run("Versión 1.0 • Septiembre 2026\nSan Salvador de Jujuy, Argentina")
    r_meta3.font.name = "Arial"
    r_meta3.font.size = Pt(9.5)
    r_meta3.font.color.rgb = COLOR_SLATE

    doc.add_page_break()


# ==============================================================================
# 1. MANUAL DE USUARIO: GOMEZ, RAMOS (USUARIOS Y SERVICIOS)
# ==============================================================================
def generar_manual_gomez_ramos(ruta_salida):
    doc = docx.Document()
    setup_page_layout(doc, "Manual de Usuario: Usuarios y Servicios")
    add_cover_page(
        doc,
        main_title="MANUAL DE USUARIO\nGESTIÓN DE USUARIOS Y SERVICIOS",
        subtitle="Guía práctica paso a paso para la administración de personal, control diario de atenciones, tarifas personalizadas, ficha de clientas y consentimiento informado.",
        authors="Gómez, Ramos",
        module_badge="Módulo de Usuarios y Servicios",
    )

    # ── CAPÍTULO 1 ──
    format_heading_1(doc, "1. Introducción y Bienvenida al Sistema")
    add_body_p(
        doc,
        "¡Bienvenidas al Sistema de Gestión Integral de Peluquería Lorena! Este manual ha sido redactado con un lenguaje simple, visual y cotidiano, pensado especialmente para que cualquier persona del salón (tanto Lorena como las estilistas y colaboradoras), sin necesidad de conocimientos avanzados en computación, pueda utilizar el sistema diario con total tranquilidad, rapidez y confianza."
    )
    add_body_p(
        doc,
        "El objetivo del sistema es modernizar la peluquería, dejando atrás las anotaciones en cuadernos que suelen perderse o mancharse con tinturas. Gracias a esta herramienta digital, el salón puede llevar el registro exacto de cada trabajo realizado, cuidar los tiempos de atención, resguardarse legalmente ante tratamientos químicos delicados y brindar un servicio personalizado y de excelencia a cada clienta."
    )
    add_callout_box(
        doc,
        "Requisitos Mínimos para el Salón",
        "Para usar el sistema sólo se necesita encender la computadora del salón, abrir el navegador web de preferencia (como Google Chrome o Microsoft Edge) e ingresar a la dirección del sistema. No hace falta instalar ningún programa complejo adicional.",
        box_type="info",
    )

    # ── CAPÍTULO 2 ──
    format_heading_1(doc, "2. Módulo de Usuarios, Seguridad y Roles")
    add_body_p(
        doc,
        "La computadora del salón es un equipo compartido que utilizan varias personas a lo largo de la jornada laboral. Por este motivo, el sistema cuenta con perfiles individuales que permiten identificar con precisión quién realizó cada servicio y resguardar la información económica y de administración exclusiva de Lorena."
    )

    format_heading_2(doc, "2.1. Los Dos Roles de Trabajo en el Salón")
    add_body_p(
        doc,
        "El sistema distingue claramente dos tipos de usuarios según las responsabilidades que tienen en el local:"
    )

    add_styled_table(
        doc,
        headers=["Rol", "Quién lo Utiliza", "Funciones Principales y Permisos"],
        data_rows=[
            [
                "Administradora",
                "Lorena (Dueña del Salón)",
                "• Control total del sistema y recaudación en tiempo real.\n• Agregar, editar o dar de baja personal de la peluquería.\n• Modificar la lista oficial de precios base del catálogo.\n• Realizar el cierre masivo de insumos y auditoría general.",
            ],
            [
                "Empleada",
                "Estilistas, Coloristas y Asistentes",
                "• Iniciar sesión personal para asentar sus propios servicios.\n• Registrar la clienta atendida y pactar precio y duración real.\n• Consultar el tarifario oficial de referencia.\n• Confeccionar consentimientos informados para decoloración/alisado.\n• Acceder a la agenda técnica y evolución de clientas.",
            ],
        ],
        col_widths=[1.4, 1.8, 3.3],
    )

    format_heading_2(doc, "2.2. Cómo Iniciar Sesión en el Salón")
    add_body_p(
        doc,
        "Al comenzar el turno, cada profesional debe ingresar con sus credenciales personales:"
    )
    add_step_p(doc, 1, "Abrir la pantalla de ingreso", "En la pantalla principal, presione el botón 'Iniciar Sesión'.")
    add_step_p(doc, 2, "Ingresar datos de acceso", "Escriba su Correo Electrónico registrado y su Contraseña en los casilleros correspondientes.")
    add_step_p(doc, 3, "Entrar al sistema", "Haga clic en el botón dorado 'Ingresar'. El sistema abrirá el panel general mostrando su nombre en la esquina superior derecha.")

    format_heading_2(doc, "2.3. Registro de Nuevos Usuarios y el PIN de Administradora (1234)")
    add_body_p(
        doc,
        "Cuando una nueva colaboradora se incorpora al equipo, debe registrarse desde la pantalla 'Registrarse'."
    )
    add_step_p(doc, 1, "Completar la ficha de registro", "Escriba Nombre Completo, Correo Electrónico, Teléfono y cree una Contraseña segura.")
    add_step_p(doc, 2, "Seleccionar el rol", "Por defecto, el sistema asigna el rol de 'Empleada'. Si quien se registra es Lorena (para una cuenta administradora), debe seleccionar la opción 'Administradora'.")
    add_step_p(
        doc,
        3,
        "Validación con PIN de Seguridad",
        "Para evitar que una empleada pueda asignarse permisos de administradora por error o sin autorización, el sistema solicitará obligatoriamente el PIN Especial de Administradora. El PIN configurado en el sistema es 1234. Si el PIN ingresado es correcto, la cuenta quedará habilitada con privilegios de gestión total."
    )

    add_callout_box(
        doc,
        "El PIN de Administradora es Confidencial",
        "El PIN 1234 sólo debe ser conocido por Lorena. Si una empleada se registra, simplemente debe dejar el rol de 'Empleada' y el sistema NO le solicitará ningún PIN, permitiéndole trabajar normalmente de inmediato.",
        box_type="warning",
    )

    format_heading_2(doc, "2.4. Gestión del Personal (Exclusivo Administradora)")
    add_body_p(
        doc,
        "Lorena dispone de una sección especial llamada 'Gestión de Personal' dentro del menú lateral. En esta pantalla puede:"
    )
    add_bullet_p(doc, "Ver el listado completo", "Revisar todos los profesionales activos en la peluquería, sus teléfonos y roles.")
    add_bullet_p(doc, "Editar datos", "Corregir teléfonos o nombres mal cargados.")
    add_bullet_p(doc, "Cambio de Contraseña", "Si una estilista olvidó su contraseña, la administradora puede restablecerla al instante desde esta sección sin perder su historial.")
    add_bullet_p(doc, "Baja de Personal (Seguridad)", "Si una colaboradora ya no trabaja en el salón, se presiona el botón 'Desactivar'. El sistema da de baja su acceso para que no pueda entrar más, pero CONSERVA todos los trabajos y servicios que realizó en el pasado para que la contabilidad y estadísticas nunca se alteren.")

    # ── CAPÍTULO 3 ──
    format_heading_1(doc, "3. Catálogo Oficial de Servicios y Tarifas Base")
    add_body_p(
        doc,
        "Peluquería Lorena cuenta con una lista de precios oficial clasificada por categorías (Cortes, Color, Mechas, Tratamientos Capilares y Otros). En el sistema, cualquier miembro del equipo puede hacer clic en 'Catálogo Servicios' para consultar los valores actualizados."
    )

    format_heading_2(doc, "3.1. El Concepto de Tarifas 'Desde'")
    add_body_p(
        doc,
        "En peluquería, muchos trabajos no tienen un precio rígido porque dependen de la cantidad de producto que demande el cabello (largo, volumen y grosor). Por eso, el catálogo establece 'Precios Base' de referencia:"
    )
    add_bullet_p(doc, "Cortes", "Damas ($25.000 / 45 min), Caballeros ($20.000 / 30 min), Niños ($20.000 / 30 min).")
    add_bullet_p(doc, "Color", "Raíz desde ($50.000 / 90 min), Shampoo Color ($35.000 / 45 min), Color Completo desde ($65.000 / 120 min).")
    add_bullet_p(doc, "Mechas y Aclaraciones", "Iluminación desde ($90.000), Reflejos desde ($95.000), Balayage desde ($120.000), Mechas Localizadas desde ($120.000).")
    add_bullet_p(doc, "Tratamientos Capilares", "Alisado desde ($80.000 / 150 min), Ampollas Nutritivas ($30.000 / 30 min), Hidratación Profunda ($25.000 / 45 min).")
    add_bullet_p(doc, "Otros Servicios", "Brushing o Planchita ($35.000), Peinados de Fiesta (desde $40.000).")

    add_callout_box(
        doc,
        "Edición de Precios Base",
        "Solamente Lorena (rol Administradora) puede modificar los precios del catálogo general. Si las listas de los proveedores aumentan, Lorena entra a 'Editar Servicio', cambia el precio base y desde ese momento todo el salón visualiza el nuevo valor.",
        box_type="tip",
    )

    # ── CAPÍTULO 4 ──
    format_heading_1(doc, "4. Control Diario de Atenciones y Personalización por Clienta")
    add_body_p(
        doc,
        "La pantalla de 'Control Diario' es el corazón operativo del salón. Funciona como una bitácora digital en vivo que muestra todo lo que está sucediendo en la peluquería durante el día:"
    )
    add_bullet_p(doc, "Horario de inicio", "A qué hora se comenzó a atender a la clienta.")
    add_bullet_p(doc, "Profesional a cargo", "Qué estilista realizó el trabajo.")
    add_bullet_p(doc, "Precio acordado y duración real", "Cuánto se cobró y cuánto tiempo llevó el servicio.")
    add_bullet_p(doc, "Recaudación en tiempo real", "Un panel superior calcula automáticamente el total de dinero producido y la cantidad de servicios completados en la fecha.")

    format_heading_2(doc, "4.1. Cómo Asentar un Servicio Realizado Paso a Paso")
    add_body_p(
        doc,
        "Cada vez que una clienta se sienta en el sillón y se acuerda el trabajo a realizar, la estilista registra la atención:"
    )
    add_step_p(doc, 1, "Abrir el formulario de registro", "En el menú lateral haga clic en 'Control Diario' y luego en el botón dorado '+ Nueva Atención'.")
    add_step_p(
        doc,
        2,
        "Seleccionar el Servicio",
        "Elija el servicio en la lista desplegable (ejemplo: 'Balayage'). En ese mismo instante, el sistema autocompleta automáticamente el casillero de precio ($120.000) y de duración (180 min)."
    )
    add_step_p(doc, 3, "Elegir a la Profesional", "Seleccione su nombre en la lista desplegable de profesionales.")
    add_step_p(doc, 4, "Datos de la Clienta", "Si la clienta ya está en la base de datos, elíjala del menú de clientas agendadas; el sistema cargará su nombre y teléfono automáticamente. Si es una clienta nueva, escriba su nombre directamente en el casillero.")
    add_step_p(
        doc,
        5,
        "Personalización Libre de Tarifa y Duración",
        "¡Este paso es clave! Supongamos que la clienta tiene cabello hasta la cintura y muy abundante. El Balayage demandará 4 horas en lugar de 3, y consumirá el doble de decolorante. La estilista simplemente borra el monto sugerido de $120.000 y escribe $150.000, y en duración cambia 180 por 240 minutos. Esto NO altera el catálogo general de Lorena, sino que registra lo que realmente se pactó con esa clienta específica."
    )
    add_step_p(doc, 6, "Guardar la Atención", "Haga clic en 'Asentar Servicio'. La atención aparecerá inmediatamente en la grilla del día.")

    # ── CAPÍTULO 5 ──
    format_heading_1(doc, "5. Ficha de Consentimiento Informado (Decoloración y Alisado)")
    add_body_p(
        doc,
        "Los procesos químicos como la decoloración extrema (rubios platinados, balayage) y los alisados térmicos o progresivos implican transformaciones profundas en la estructura del cabello. Existen múltiples factores previos y biológicos totalmente ajenos a la peluquera que pueden provocar una rotura capilar (corte químico) o reacciones alérgicas si no se conocen de antemano."
    )
    add_body_p(
        doc,
        "Por este motivo, el sistema incluye una Ficha de Consentimiento Informado con valor técnico y legal, diseñada para proteger la reputación profesional de Lorena y de sus estilistas."
    )

    format_heading_2(doc, "5.1. Factores Ajenos que la Clienta Debe Declarar Obligatoriamente")
    add_body_p(
        doc,
        "Antes de aplicar peróxidos o ácidos, la clienta debe responder con sinceridad el cuestionario guiado del sistema:"
    )
    add_bullet_p(
        doc,
        "Uso previo de Henna o Sales Metálicas",
        "¡Riesgo crítico! Los tintes vegetales comerciales o caseros suelen contener sales metálicas. Si entran en contacto con el polvo decolorante y oxidante, provocan una reacción exotérmica violenta (el cabello humea, levanta temperatura extrema y se corta de raíz)."
    )
    add_bullet_p(
        doc,
        "Alisados o Permanentes Previos",
        "Incompatibilidad química severa entre hidróxidos y tioglicolatos con decoloraciones superpuestas."
    )
    add_bullet_p(
        doc,
        "Decoloraciones previas en largos y puntas",
        "Cabello con porosidad extrema que requiere un tratamiento de nutrición previo antes de volver a aclarar."
    )
    add_bullet_p(
        doc,
        "Alergias o Cuero Cabelludo Sensible",
        "Antecedentes de dermatitis, ardor o picazón ante persulfatos o amoníaco."
    )
    add_bullet_p(
        doc,
        "Embarazo, Lactancia o Medicación Fuerte",
        "Cambios hormonales o fármacos (como tratamientos tiroideos o quimioterapia) que debilitan la raíz del folículo."
    )

    format_heading_2(doc, "5.2. Diagnóstico Técnico de la Profesional")
    add_body_p(
        doc,
        "En la misma ficha, la estilista asienta el resultado de la Prueba de Mecha obligatoria (Apto / Apto con Precaución / No Apto) y el estado de elasticidad y porosidad de la fibra capilar."
    )

    format_heading_2(doc, "5.3. Impresión y Firma de la Ficha")
    add_body_p(
        doc,
        "Al guardar la ficha, el sistema genera automáticamente un documento formal con texto legal de deslinde de responsabilidad por omisión de datos. La profesional presiona el botón 'Imprimir Ficha' para imprimirla en papel y hacerla firmar de puño y letra por la clienta, o bien guardarla digitalmente en PDF."
    )

    # ── CAPÍTULO 6 ──
    format_heading_1(doc, "6. Agenda de Clientas y Seguimiento Multisesión")
    add_body_p(
        doc,
        "En la sección 'Agenda Clientas', el salón cuenta con una base de datos organizada para fidelizar al público y llevar un registro detallado de su evolución."
    )

    format_heading_2(doc, "6.1. WhatsApp Directo y Alerta de Cumpleaños")
    add_bullet_p(
        doc,
        "Botón de WhatsApp en un clic",
        "Al lado de cada clienta hay un botón verde de WhatsApp. Al tocarlo en la computadora, se abre de inmediato el chat oficial con la clienta sin necesidad de agendar el número en el teléfono personal."
    )
    add_bullet_p(
        doc,
        "Detector de Cumpleaños para Promociones",
        "En la parte superior de la pantalla, el sistema coloca un cartel festivo destacado indicando qué clientas cumplen años hoy o en los próximos 15 días. Esto permite enviarles con un solo toque un mensaje de felicitaciones, un descuento especial de cumpleaños o un regalo de hidratación."
    )

    format_heading_2(doc, "6.2. Seguimiento Técnico de Tratamientos Multisesión")
    add_body_p(
        doc,
        "Ciertos objetivos (por ejemplo, pasar de un cabello teñido de negro a un rubio ceniza, o un alisado progresivo en varias aplicaciones) demandan 2, 3 o más citas espaciadas en semanas para no quebrar el cabello. Para esto existe la Ficha Multisesión:"
    )
    add_step_p(doc, 1, "Iniciar Tratamiento", "Entre a la ficha de la clienta y presione 'Iniciar Tratamiento'. Ingrese el título (ej: 'Transición Balayage Platinado') y la cantidad de sesiones estimadas (ej: 3 sesiones).")
    add_step_p(
        doc,
        2,
        "Registrar cada Sesión Técnica",
        "En cada visita, presione 'Registrar Sesión'. La estilista anota con precisión técnica:\n• Diagnóstico de la fibra en ese día.\n• Fórmulas químicas utilizadas (ej: 'Polvo Decolorante + Oxidante 20 vol + Plex protector + Matizador 9.12').\n• Tiempo de exposición (ej: '45 minutos').\n• Resultado obtenido.\n• Recomendaciones de cuidado para el hogar (ej: 'Usar shampoo libre de sulfatos y mascarilla de nutrición 2 veces por semana')."
    )
    add_step_p(
        doc,
        3,
        "Barra de Progreso Automática",
        "El sistema calcula el porcentaje de avance (33%, 66%, 100%) y muestra el historial cronológico para que, si en la siguiente cita a la clienta la atiende otra compañera, sepa con exactitud qué fórmula se le aplicó la última vez."
    )

    # ── CAPÍTULO 7 ──
    format_heading_1(doc, "7. Cierre Diario de Servicios y Descuento Masivo de Insumos")
    add_body_p(
        doc,
        "Al finalizar el día laboral, el equipo no tiene que estar restando a mano producto por producto. Se utiliza el panel de 'Cierre de Insumos'."
    )

    format_heading_2(doc, "7.1. Cómputo Automático de Personas por Servicio")
    add_body_p(
        doc,
        "Al ingresar a la pantalla de cierre con la fecha de hoy, el sistema agrupa y cuenta cuántas personas se hicieron el mismo servicio a lo largo de toda la jornada. Por ejemplo:"
    )
    add_bullet_p(doc, "Balayage", "4 personas atendidas en el día.")
    add_bullet_p(doc, "Cortes Damas", "6 personas atendidas en el día.")
    add_bullet_p(doc, "Alisados", "2 personas atendidas en el día.")

    format_heading_2(doc, "7.2. Descuento Masivo de Stock en Lote")
    add_step_p(
        doc,
        1,
        "Cargar los insumos consumidos",
        "Para el grupo de 'Balayage' (4 personas), la responsable mira la tarjeta y selecciona los insumos gastados en total: por ejemplo, selecciona 'Polvo Decolorante 500g' y coloca cantidad: 2 potes; luego selecciona 'Oxigenta 20 vol' y coloca cantidad: 3 botellas."
    )
    add_step_p(
        doc,
        2,
        "Presionar 'Descontar Stock'",
        "Al presionar el botón, el sistema descuenta automáticamente esas cantidades del inventario general en un solo paso, genera la auditoría del movimiento y marca ese servicio del día como cerrado."
    )
    add_step_p(
        doc,
        3,
        "Protección contra Dobles Cierres",
        "Si alguien intenta volver a presionar el botón de cierre para ese mismo servicio en la misma fecha, el sistema emitirá un aviso y bloqueará la operación, evitando que el stock se descuente dos veces por error."
    )

    # ── CAPÍTULO 8 ──
    format_heading_1(doc, "8. Preguntas Frecuentes y Consejos Prácticos")
    add_bullet_p(
        doc,
        "¿Qué pasa si me equivoqué al cargar el precio de una atención?",
        "Mientras el servicio esté en la bitácora del día, la administradora puede editar la atención o cancelarla para volver a cargar el monto correcto."
    )
    add_bullet_p(
        doc,
        "¿Es obligatorio hacer el consentimiento informado para un simple corte?",
        "No. El consentimiento sólo es requerido para procesos químicos que alteran la estructura del cabello (decoloraciones, mechas y alisados)."
    )
    add_bullet_p(
        doc,
        "¿Por qué es importante cerrar sesión al terminar el día?",
        "Para que la computadora quede lista para la mañana siguiente y nadie pueda realizar operaciones con el usuario de otra persona."
    )

    doc.save(ruta_salida)
    print(f"[OK] Generado exitosamente: {ruta_salida}")


# ==============================================================================
# 2. MANUAL DE USUARIO: MANSILLA, CHIACCHIO (INVENTARIOS Y PROVEEDORES)
# ==============================================================================
def generar_manual_mansilla_chiacchio(ruta_salida):
    doc = docx.Document()
    setup_page_layout(doc, "Manual de Usuario: Inventarios y Proveedores")
    add_cover_page(
        doc,
        main_title="MANUAL DE USUARIO\nCONTROL DE INVENTARIOS Y PROVEEDORES",
        subtitle="Guía práctica paso a paso para el control de stock, reposición de insumos, alertas de stock mínimo, gestión de distribuidores y recepción de pedidos.",
        authors="Mansilla, Chiacchio",
        module_badge="Módulo de Inventarios y Proveedores",
    )

    # ── CAPÍTULO 1 ──
    format_heading_1(doc, "1. Introducción al Control de Stock del Salón")
    add_body_p(
        doc,
        "¡Bienvenidas al Manual de Inventarios y Proveedores de Peluquería Lorena! En un salón de belleza profesional, los productos e insumos representan el corazón del negocio y una inversión económica fundamental. Quedarse sin polvo decolorante o sin tintura castaño oscuro un sábado por la tarde significa perder clientas y dinero; y comprar mercadería de más por no saber qué hay en el depósito inmoviliza capital innecesariamente."
    )
    add_body_p(
        doc,
        "Este manual ha sido elaborado con un lenguaje claro, accesible y sin tecnicismos informáticos, para que todo el equipo del salón pueda mantener el inventario ordenado, conocer en tiempo real cuántas unidades quedan de cada producto y realizar pedidos a los distribuidores de manera fácil y segura."
    )

    format_heading_2(doc, "1.1. Los Tipos de Productos que Maneja el Salón")
    add_body_p(
        doc,
        "En el sistema de Peluquería Lorena se gestionan principalmente dos clases de mercadería:"
    )
    add_bullet_p(
        doc,
        "Insumos Técnicos de Uso Interno",
        "Productos que utiliza el equipo en la bacha y en los sillones para realizar los servicios: polvos decolorantes, oxidantes de diferentes volúmenes (10, 20, 30, 40 vol.), tinturas profesionales, ampollas de nutrición, líquidos de alisado y botox capilar."
    )
    add_bullet_p(
        doc,
        "Productos para Reventa al Público",
        "Líneas de cuidado profesional que las clientas compran en el mostrador para continuar el tratamiento en su hogar: champús matizadores, mascarillas nutritivas, aceites de argán y protectores térmicos."
    )

    # ── CAPÍTULO 2 ──
    format_heading_1(doc, "2. Consulta del Stock y Catálogo de Insumos")
    add_body_p(
        doc,
        "Para revisar el estado de los insumos, haga clic en la opción 'Inventario' del menú lateral. En pantalla se desplegará la tabla completa de productos registrados en el salón."
    )

    format_heading_2(doc, "2.1. Cómo Interpretar la Tabla de Productos")
    add_body_p(
        doc,
        "Cada fila de la tabla representa un producto y contiene las siguientes columnas informativas:"
    )

    add_styled_table(
        doc,
        headers=["Dato / Columna", "Qué Significa", "Para qué Sirve"],
        data_rows=[
            ["Producto", "Nombre comercial y presentación (ej: 'Polvo Decolorante 500g').", "Identificar rápidamente el insumo."],
            ["Stock Actual", "Número de unidades físicas que deben estar en los estantes del salón.", "Saber cuántos frascos o pomos quedan disponibles hoy."],
            ["Stock Mínimo", "Cantidad de seguridad que nunca debería faltar en el local.", "El límite de advertencia para saber cuándo pedir más."],
            ["Estado / Alerta", "Insignia de color verde, amarilla o roja.", "Avisa de un vistazo si el producto está en nivel seguro o crítico."],
            ["Precio", "Precio de costo o sugerido del producto.", "Valorizar el inventario del salón."],
            ["Acciones", "Botones '+ Reponer' y '- Descontar'.", "Registrar entradas y salidas de stock."],
        ],
        col_widths=[1.5, 2.5, 2.5],
    )

    format_heading_2(doc, "2.2. Buscador Rápido de Productos")
    add_body_p(
        doc,
        "Si la peluquería cuenta con decenas de tonos de tinturas o marcas distintas, no hace falta buscar una por una en la lista. En la parte superior de la tabla hay un Buscador Inteligente: simplemente escriba una palabra (por ejemplo: 'Oxigenta' o 'Alisado') y presione Enter. La pantalla filtrará en el acto los productos que coincidan con esa palabra."
    )

    # ── CAPÍTULO 3 ──
    format_heading_1(doc, "3. Sistema de Alertas y Semáforo de Stock Mínimo")
    add_body_p(
        doc,
        "El mayor problema de una peluquería es descubrir que se terminó un producto justo en el momento en que una clienta tiene el decolorante puesto en la cabeza. Para evitar esta situación, el sistema cuenta con un Semáforo Automático de Stock Mínimo."
    )

    format_heading_2(doc, "3.1. ¿Cómo Funciona el Stock Mínimo?")
    add_body_p(
        doc,
        "Supongamos que para la tintura 'Tono 7.1 Rubio Ceniza' Lorena definió un Stock Mínimo de 3 pomos. Esto significa que siempre deben haber al menos 3 unidades de reserva en el estante para trabajar con tranquilidad:"
    )
    add_bullet_p(doc, "🟢 Nivel Seguro (Verde)", "Si hay 6 pomos en stock, el número se muestra en verde con la leyenda 'Stock Normal'. Hay suficiente mercadería.")
    add_bullet_p(doc, "🟡 Nivel de Advertencia (Amarillo)", "Si quedan exactamente 3 pomos (se alcanzó el límite mínimo), el sistema muestra una advertencia amarilla indicando que es momento de incluir este producto en el próximo pedido al distribuidor.")
    add_bullet_p(doc, "🔴 Nivel Crítico (Rojo)", "Si quedan 1 o 2 pomos, o si el stock llega a 0, la insignia se vuelve roja con la leyenda '¡Stock Crítico!'. Se debe pedir reposición urgente.")

    format_heading_2(doc, "3.2. Pestaña de 'Productos con Bajo Stock'")
    add_body_p(
        doc,
        "Cuando Lorena o la encargada de compras va a llamar al viajante de la distribuidora, no necesita revisar estante por estante. Simplemente hace clic en el botón 'Ver Bajo Stock' y el sistema le muestra en una sola pantalla únicamente los productos que alcanzaron o perforaron su stock mínimo, listos para armar la orden de compra en 2 minutos."
    )

    # ── CAPÍTULO 4 ──
    format_heading_1(doc, "4. Salidas y Descuento Manual de Insumos")
    add_body_p(
        doc,
        "El descuento principal de insumos del salón se realiza al terminar el día desde el módulo de servicios (descuento masivo por clientas atendidas). Sin embargo, en el día a día de un local ocurren situaciones imprevistas donde es necesario restar unidades a mano."
    )

    format_heading_2(doc, "4.1. ¿Cuándo se debe usar el Descuento Manual?")
    add_bullet_p(doc, "Rotura Accidental", "Un pomo de tintura o una botella de oxidante se cayó en el lavacabezas o en el piso y se rompió.")
    add_bullet_p(doc, "Producto Vencido o Deteriorado", "Un producto que estuvo mucho tiempo abierto perdió su efectividad y debe descartarse.")
    add_bullet_p(doc, "Venta Directa de Mostrador", "Una clienta compró un champú de reventa para llevarse a su casa.")
    add_bullet_p(doc, "Uso Interno Extraordinario", "Limpieza profunda de toallas o desinfección de herramientas con productos de stock.")

    format_heading_2(doc, "4.2. Paso a Paso para Descontar Stock")
    add_step_p(doc, 1, "Ubicar el Insumo", "Busque el producto en la lista de inventario.")
    add_step_p(doc, 2, "Hacer clic en '- Descontar'", "Presione el botón rojo '- Descontar' ubicado en la columna de acciones.")
    add_step_p(doc, 3, "Indicar la Cantidad", "Escriba cuántas unidades se retiran del salón (ejemplo: 1 pomo).")
    add_step_p(doc, 4, "Escribir el Motivo Obligatorio", "Escriba con claridad por qué sale ese producto (ejemplo: 'Rotura de envase en la bacha' o 'Venta en mostrador a clienta María').")
    add_step_p(doc, 5, "Confirmar la Operación", "Presione 'Confirmar Descuento'. El stock disminuirá en el acto y el motivo quedará guardado para siempre en la auditoría.")

    add_callout_box(
        doc,
        "Protección contra Stock Negativo",
        "El sistema es inteligente y nunca permitirá descontar más unidades de las que realmente existen. Si en el salón hay 2 botellas de champú y alguien intenta descontar 5, el sistema rechazará la operación con un cartel de error, evitando que los números de stock queden en negativo.",
        box_type="danger",
    )

    # ── CAPÍTULO 5 ──
    format_heading_1(doc, "5. Entrada y Reposición Manual de Mercadería")
    add_body_p(
        doc,
        "Cuando ingresa mercadería nueva al salón que no proviene de una orden de compra formal (por ejemplo, si Lorena compró 3 frascos de crema de urgencia en una distribuidora del centro con dinero de la caja chica), se realiza una reposición manual rápida."
    )

    format_heading_2(doc, "5.1. Paso a Paso para Sumar Stock")
    add_step_p(doc, 1, "Ubicar el Insumo", "En la pantalla de inventario, busque el producto que acaba de llegar.")
    add_step_p(doc, 2, "Hacer clic en '+ Reponer'", "Presione el botón verde '+ Reponer'.")
    add_step_p(doc, 3, "Ingresar las Unidades Nuevas", "Escriba la cantidad de unidades que entraron al salón (ejemplo: 6 botellas).")
    add_step_p(doc, 4, "Indicar el Motivo o Factura", "Escriba el comprobante o referencia (ejemplo: 'Compra urgente en Distribuidora Norte - Ticket #8492').")
    add_step_p(doc, 5, "Confirmar la Entrada", "Presione 'Confirmar Reposición'. Las unidades se sumarán al instante al stock actual.")

    # ── CAPÍTULO 6 ──
    format_heading_1(doc, "6. Historial y Auditoría de Movimientos")
    add_body_p(
        doc,
        "¿Alguna vez faltó un producto en el salón y nadie sabía qué había pasado con él? Para terminar con los misterios y garantizar la máxima transparencia en el equipo, el sistema cuenta con un Historial de Movimientos que audita cada movimiento de stock."
    )

    format_heading_2(doc, "6.1. Información Registrada en Cada Movimiento")
    add_body_p(
        doc,
        "Cada vez que alguien suma o resta una unidad, el sistema guarda de forma inalterable:"
    )
    add_bullet_p(doc, "Fecha y Hora exacta", "El momento preciso en que se hizo la operación.")
    add_bullet_p(doc, "Responsable", "El nombre de la persona que inició sesión y tocó el stock.")
    add_bullet_p(doc, "Tipo de Movimiento", "Si fue una 'Entrada' (verde) o una 'Salida' (roja).")
    add_bullet_p(doc, "Cantidad", "El número de unidades modificadas.")
    add_bullet_p(doc, "Motivo Registrado", "La justificación escrita de la operación.")

    add_callout_box(
        doc,
        "Control para la Dueña del Salón",
        "Lorena puede consultar este historial filtrando por fechas para saber exactamente cuántos pomos de tintura se gastaron en la semana y verificar que el consumo físico coincida con los trabajos cobrados.",
        box_type="tip",
    )

    # ── CAPÍTULO 7 ──
    format_heading_1(doc, "7. Directorio y Gestión de Proveedores")
    add_body_p(
        doc,
        "En la peluquería se trabaja con distintos distribuidores de cosmética capilar (L'Oréal, Wella, Silkey, Nov, etc.). Para no depender de tener los números de teléfono anotados en papeles sueltos o en el celular privado de una sola persona, el sistema incluye la agenda central de 'Proveedores'."
    )

    format_heading_2(doc, "7.1. Cómo Dar de Alta a un Proveedor")
    add_step_p(doc, 1, "Ingresar a Proveedores", "En el menú lateral haga clic en 'Proveedores' y luego en '+ Nuevo Proveedor'.")
    add_step_p(doc, 2, "Datos de la Empresa", "Escriba la Razón Social o Nombre de la Distribuidora.")
    add_step_p(doc, 3, "Contacto Directo", "Coloque el nombre del viajante o vendedor que atiende el local y su número de WhatsApp / Teléfono.")
    add_step_p(doc, 4, "Datos de Entrega y Pago", "Agregue dirección, correo electrónico y notas útiles (ejemplo: 'Viene los martes por la mañana; acepta transferencia a 15 días').")
    add_step_p(doc, 5, "Guardar", "El proveedor quedará agendado y disponible para emitirle pedidos en el sistema.")

    # ── CAPÍTULO 8 ──
    format_heading_1(doc, "8. Órdenes de Compra y Recepción de Pedidos")
    add_body_p(
        doc,
        "Este módulo permite cerrar el circuito de compras de manera sumamente profesional y organizada, evitando errores en la entrega de mercadería."
    )

    format_heading_2(doc, "8.1. El Circuito de Pedidos en 4 Pasos")
    add_step_p(
        doc,
        1,
        "Crear la Orden de Compra",
        "Haga clic en 'Pedidos a Proveedores' y presione '+ Nuevo Pedido'. Elija el proveedor al que se le va a comprar (ej: 'Distribuidora Bella Salón')."
    )
    add_step_p(
        doc,
        2,
        "Seleccionar los Productos y Cantidades",
        "Agregue a la lista los insumos que están haciendo falta (apoyándose en la lista de bajo stock) e indique la cantidad de unidades solicitadas y el costo pactado."
    )
    add_step_p(
        doc,
        3,
        "Enviar el Pedido",
        "El sistema genera la orden con estado 'Pendiente de Entrega'. El salón puede descargar o copiar el detalle para mandárselo por WhatsApp al vendedor."
    )
    add_step_p(
        doc,
        4,
        "Recepción Automática de Mercadería",
        "¡La mayor ventaja del sistema! Cuando el repartidor llega al local con las cajas, la responsable abre el pedido en la computadora, coteja que las cantidades coincidan con la factura y presiona el botón 'Confirmar Recepción'.\nAl presionar este único botón, el sistema CAMBIA el estado del pedido a 'Recibido' y SUMA AUTOMÁTICAMENTE todas las unidades al inventario de cada producto, sin necesidad de cargarlas una por una a mano."
    )

    # ── CAPÍTULO 9 ──
    format_heading_1(doc, "9. Consejos de Oro para el Cuidado del Stock")
    add_bullet_p(doc, "Registrar las roturas en el momento", "Si un producto se derrama o se cae, no lo deje para después: cargue el descuento en el sistema de inmediato para que los números no se desfasen.")
    add_bullet_p(doc, "Control de vencimientos", "Aplique el principio 'lo primero que entra es lo primero que sale': coloque los productos nuevos detrás de los que ya estaban abiertos en el estante.")
    add_bullet_p(doc, "Inventario físico mensual", "Una vez al mes, tómense 20 minutos con Lorena para contar los estantes y verificar que coincidan al 100% con la tabla del sistema.")

    doc.save(ruta_salida)
    print(f"[OK] Generado exitosamente: {ruta_salida}")


if __name__ == "__main__":
    carpeta = r"y:\peluqueria_lorena"
    ruta_gomez_ramos = os.path.join(carpeta, "Manual de Usuario- Gomez, Ramos.docx")
    ruta_mansilla_chiacchio = os.path.join(carpeta, "Manual de Usuario- Mansilla, Chiacchio.docx")

    print("Generando Manual de Usuario - Gomez, Ramos...")
    generar_manual_gomez_ramos(ruta_gomez_ramos)

    print("Generando Manual de Usuario - Mansilla, Chiacchio...")
    generar_manual_mansilla_chiacchio(ruta_mansilla_chiacchio)

    print("\n¡Ambos manuales han sido generados exitosamente!")
